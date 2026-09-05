from pathlib import Path

try:
    import pymupdf as fitz
except ImportError:
    import fitz


ALLOWED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".txt"
}


def _render_table(table) -> str:
    """Render a detected PDF table as pipe-delimited rows.

    The relationship between columns (e.g. Date -> Festival -> Day)
    is preserved so retrieval and chunking keep rows intact.
    """
    try:
        data = table.extract()
    except Exception:
        return ""

    if not data:
        return ""

    rows = []
    for row in data:
        cells = []
        for cell in row:
            value = str(cell).replace("|", "/").replace("\n", " ").strip() if cell else ""
            if value == "None":
                value = ""
            cells.append(value)
        rows.append(" | ".join(cells))

    return "\n".join(rows)


def _extract_pdf_tables(pdf, page_number, page) -> str:
    """Return pipe-delimited tables found on a PDF page, if any."""
    find_tables = getattr(page, "find_tables", None)
    if find_tables is None:
        return ""

    try:
        tables = find_tables()
    except Exception:
        return ""

    parts = []
    try:
        for table in tables:
            rendered = _render_table(table)
            if rendered:
                parts.append(rendered)
    except Exception:
        pass

    return "\n\n".join(parts)


def load_pdf(file_path: str):
    pages = []
    try:
        pdf = fitz.open(file_path)
    except Exception as err:
        raise ValueError(f"Failed to open PDF file: {str(err)}")

    try:
        for page_number, page in enumerate(pdf, start=1):
            text = page.get_text("text").strip()
            table_text = _extract_pdf_tables(pdf, page_number, page)
            if table_text:
                text = f"{text}\n\nTable:\n{table_text}" if text else table_text
            if text.strip():
                pages.append({
                    "text": text.strip(),
                    "page": page_number
                })
    finally:
        pdf.close()

    return pages


def _extract_docx_tables(document) -> str:
    parts = []
    try:
        for table in document.tables:
            rows = []
            for row in table.rows:
                cells = []
                for cell in row.cells:
                    value = cell.text.replace("|", "/").replace("\n", " ").strip()
                    cells.append(value)
                rows.append(" | ".join(cells))
            if rows:
                parts.append("\n".join(rows))
    except Exception:
        pass
    return "\n\n".join(parts)


def load_docx(file_path: str):
    try:
        from docx import Document
    except ImportError:
        raise ValueError(
            "python-docx is not installed. "
            "Add 'python-docx' to requirements.txt."
        )
    try:
        document = Document(file_path)
    except Exception as err:
        raise ValueError(f"Failed to open DOCX file: {str(err)}")

    paragraphs = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            paragraphs.append(text)

    text = "\n".join(paragraphs)
    table_text = _extract_docx_tables(document)
    if table_text:
        text = f"{text}\n\nTable:\n{table_text}" if text else table_text

    if not text.strip():
        return []

    return [
        {
            "text": text.strip(),
            "page": None
        }
    ]


def load_txt(file_path: str):
    try:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            text = f.read().strip()
    except Exception as err:
        raise ValueError(f"Failed to read TXT file: {str(err)}")

    if not text:
        return []

    return [
        {
            "text": text,
            "page": None
        }
    ]


def load_document(file_path: str):
    extension = Path(file_path).suffix.lower()

    if extension == ".pdf":
        return load_pdf(file_path)

    if extension == ".docx":
        return load_docx(file_path)

    if extension == ".txt":
        return load_txt(file_path)

    raise ValueError("Only PDF, DOCX, and TXT files are supported.")